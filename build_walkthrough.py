from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = "document_intelligence_walkthrough.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(16, 42, 42)
TEAL = RGBColor(8, 127, 115)
MINT = RGBColor(231, 242, 238)
INK = RGBColor(36, 53, 51)
MUTED = RGBColor(83, 100, 97)
PAPER = RGBColor(246, 248, 245)
WHITE = RGBColor(255, 255, 255)
CORAL = RGBColor(220, 104, 76)
GOLD = RGBColor(224, 171, 68)


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False, font="Aptos", align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(7)
        if align:
            p.alignment = align
    return shape


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def base(title, kicker, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    box(slide, 0, 0, 13.333, 0.16, TEAL)
    textbox(slide, kicker.upper(), 0.7, 0.45, 6, 0.3, 10, TEAL, True)
    textbox(slide, title, 0.7, 0.78, 11.8, 0.65, 27, NAVY, True, "Aptos Display")
    textbox(slide, f"{number:02d}", 12.1, 0.48, 0.55, 0.35, 11, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def bullet_block(slide, items, x, y, w, h, size=17, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(12)
        p.level = 0
        p.text = "  " + item
    return shape

# 1 title
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
box(slide, 0, 0, 13.333, 0.2, CORAL)
textbox(slide, "DOCUMENT INTELLIGENCE", 0.8, 0.8, 7, 0.4, 13, MINT, True)
textbox(slide, "From uploaded evidence\nto actionable answers", 0.8, 1.45, 8.4, 1.7, 38, WHITE, True, "Aptos Display")
textbox(slide, "Technical walkthrough | FastAPI + Streamlit + AI services + Docker Hub + EC2", 0.85, 3.55, 9.5, 0.45, 17, MINT)
box(slide, 0.85, 5.35, 3.1, 0.08, TEAL)
textbox(slide, "A practical guide to the codebase and deployment path", 0.85, 5.6, 7.5, 0.4, 15, WHITE)
textbox(slide, "2026", 11.3, 6.65, 1.2, 0.3, 12, MINT, True, align=PP_ALIGN.RIGHT)

# 2 purpose
slide = base("What the system does", "Product scope", 2)
box(slide, 0.7, 1.8, 3.75, 4.65, NAVY, radius=True)
textbox(slide, "01", 1.05, 2.15, 0.7, 0.4, 18, CORAL, True)
textbox(slide, "Extract", 1.05, 2.7, 2.6, 0.45, 25, WHITE, True, "Aptos Display")
textbox(slide, "OCR and VLM paths turn invoices and images into searchable text and structured fields.", 1.05, 3.35, 2.85, 1.4, 16, MINT)
box(slide, 4.8, 1.8, 3.75, 4.65, TEAL, radius=True)
textbox(slide, "02", 5.15, 2.15, 0.7, 0.4, 18, GOLD, True)
textbox(slide, "Understand", 5.15, 2.7, 3.0, 0.45, 25, WHITE, True, "Aptos Display")
textbox(slide, "Voice questions can be combined with an optional image and conversation context.", 5.15, 3.35, 2.9, 1.4, 16, WHITE)
box(slide, 8.9, 1.8, 3.75, 4.65, CORAL, radius=True)
textbox(slide, "03", 9.25, 2.15, 0.7, 0.4, 18, NAVY, True)
textbox(slide, "Respond", 9.25, 2.7, 2.7, 0.45, 25, NAVY, True, "Aptos Display")
textbox(slide, "The pipeline returns a transcript, an answer, and generated audio while persisting history.", 9.25, 3.35, 2.9, 1.4, 16, NAVY)

# 3 architecture
slide = base("Architecture at a glance", "System map", 3)
labels = [("Browser", 0.9, 2.4, NAVY), ("Streamlit\nUI :8501", 3.25, 2.4, TEAL), ("FastAPI\nAPI :8000", 6.0, 2.4, CORAL), ("Services", 8.75, 1.55, GOLD), ("Data + storage", 8.75, 3.3, NAVY)]
for label, x, y, color in labels:
    box(slide, x, y, 2.0, 1.1, color, radius=True)
    textbox(slide, label, x + .15, y + .28, 1.7, .55, 17, WHITE if color != GOLD else NAVY, True, align=PP_ALIGN.CENTER)
for x, y, w in [(2.9, 2.92, .35), (5.65, 2.92, .35), (8.4, 2.0, .35), (8.4, 3.75, .35)]:
    box(slide, x, y, w, .08, MUTED)
textbox(slide, "HTTP + bearer token", 2.35, 2.45, 1.3, .3, 10, MUTED, align=PP_ALIGN.CENTER)
textbox(slide, "authenticated API", 5.0, 2.45, 1.4, .3, 10, MUTED, align=PP_ALIGN.CENTER)
textbox(slide, "OCR / VLM / STT / TTS", 7.25, 1.35, 1.4, .3, 10, MUTED, align=PP_ALIGN.CENTER)
textbox(slide, "PostgreSQL, Redis, Supabase", 7.1, 4.95, 2.8, .3, 11, MUTED)
textbox(slide, "The UI is deliberately thin: it handles user interaction and forwards files to the API. Business logic stays in the backend services.", 1.0, 5.55, 11.0, .6, 17, INK)

# 4 layout
slide = base("How the repository is organized", "Code navigation", 4)
box(slide, 0.75, 1.7, 5.9, 4.9, NAVY, radius=True)
textbox(slide, "app/", 1.1, 2.05, 2, .4, 24, MINT, True, "Aptos Display")
textbox(slide, "main.py                 application + CORS\nconfig.py               environment settings\ndatabase.py             async SQLAlchemy\nschema.py               request/response models\nmodels.py               persistence models\napi/                    HTTP route modules\ndependencies/auth.py    JWT dependency\nservices/               OCR, VLM, STT, TTS, storage\ntests/                  provider-mocked tests", 1.1, 2.75, 4.9, 2.9, 15, WHITE)
box(slide, 7.0, 1.7, 5.55, 4.9, MINT, radius=True)
textbox(slide, "Delivery files", 7.4, 2.05, 3, .4, 24, NAVY, True, "Aptos Display")
textbox(slide, "streamlit_app.py       browser-facing workspace\nDockerfile.api          FastAPI image\nDockerfile.ui           Streamlit image\ndocker-compose.yml      pull-only runtime\n.github/workflows/      Docker Hub + EC2 CI/CD\nREADME.md               setup and operations", 7.4, 2.8, 4.5, 2.3, 16, INK)

# 5 auth
slide = base("Authentication and request boundaries", "Security path", 5)
textbox(slide, "Registration and login both return a JWT.", 0.8, 1.65, 5.8, .4, 20, NAVY, True)
box(slide, 0.8, 2.35, 2.5, 1.2, MINT, radius=True)
textbox(slide, "POST /auth/login\nemail + password", 1.0, 2.68, 2.1, .55, 16, NAVY, True, align=PP_ALIGN.CENTER)
box(slide, 3.8, 2.35, 2.5, 1.2, TEAL, radius=True)
textbox(slide, "access_token\nBearer JWT", 4.0, 2.68, 2.1, .55, 16, WHITE, True, align=PP_ALIGN.CENTER)
box(slide, 6.8, 2.35, 2.5, 1.2, NAVY, radius=True)
textbox(slide, "Protected route\nget_current_user", 7.0, 2.68, 2.1, .55, 16, WHITE, True, align=PP_ALIGN.CENTER)
box(slide, 9.8, 2.35, 2.5, 1.2, CORAL, radius=True)
textbox(slide, "User-scoped data\nDB ownership checks", 10.0, 2.68, 2.1, .55, 16, NAVY, True, align=PP_ALIGN.CENTER)
for x in [3.35, 6.35, 9.35]: box(slide, x, 2.9, .35, .08, MUTED)
bullet_block(slide, ["Passwords are hashed with bcrypt.", "The auth dependency decodes the token, resolves the user, and rejects missing or invalid users.", "CORS_ORIGINS limits browser origins; credentials are disabled because the UI sends the JWT explicitly."], 0.95, 4.35, 11.3, 1.7, 16)

# 6 document flow
slide = base("Document extraction flow", "Workflow 01", 6)
steps = [("Upload", "PNG or JPEG", MINT), ("Validate", "type + size", TEAL), ("Process", "OCR or VLM", NAVY), ("Persist", "storage + DB", CORAL), ("Render", "text + fields", GOLD)]
for i, (head, sub, color) in enumerate(steps):
    x = .75 + i * 2.5
    box(slide, x, 2.0, 1.95, 1.55, color, radius=True)
    textbox(slide, f"0{i+1}", x + .15, 2.2, .4, .25, 11, WHITE if color != GOLD else NAVY, True)
    textbox(slide, head, x + .15, 2.62, 1.65, .35, 19, WHITE if color != GOLD else NAVY, True, align=PP_ALIGN.CENTER)
    textbox(slide, sub, x + .15, 3.08, 1.65, .25, 12, WHITE if color != GOLD else NAVY, align=PP_ALIGN.CENTER)
    if i < 4: box(slide, x + 2.05, 2.72, .35, .08, MUTED)
textbox(slide, "OCR endpoint: /documents/\nVLM endpoint: /documents/vlm", 1.0, 4.55, 3.5, .8, 19, NAVY, True)
textbox(slide, "Both routes require a bearer token, enforce PNG/JPEG boundaries, store the original file, and return a DocumentResponse. The OCR route additionally extracts layout and structured invoice data.", 5.0, 4.45, 7.0, 1.0, 17, INK)

# 7 multimodal
slide = base("Multimodal pipeline", "Workflow 02", 7)
box(slide, .8, 1.85, 11.7, 3.9, WHITE, line=MINT, radius=True)
flow = [("audio", "STT", TEAL), ("transcript", "context", NAVY), ("image + prompt", "VLM / LLM", CORAL), ("response", "TTS", GOLD), ("audio path", "persist", NAVY)]
for i, (head, sub, color) in enumerate(flow):
    x = 1.15 + i * 2.25
    box(slide, x, 3.0, 1.65, 1.1, color, radius=True)
    textbox(slide, head, x + .08, 3.25, 1.5, .25, 14, WHITE if color != GOLD else NAVY, True, align=PP_ALIGN.CENTER)
    textbox(slide, sub, x + .08, 3.6, 1.5, .25, 11, WHITE if color != GOLD else NAVY, align=PP_ALIGN.CENTER)
    if i < 4: box(slide, x + 1.75, 3.52, .28, .07, MUTED)
textbox(slide, "POST /multimodal/voice-image", 1.15, 2.2, 4.2, .35, 19, NAVY, True)
textbox(slide, "The API creates or reuses a conversation, loads the latest context, checks Redis, runs the async service pipeline, uploads source and generated files concurrently, commits messages, and caches the response.", 1.15, 4.7, 10.7, .7, 16, INK)

# 8 UI
slide = base("Streamlit workspace", "User experience", 8)
box(slide, .8, 1.7, 4.3, 4.9, NAVY, radius=True)
textbox(slide, "SIGN IN", 1.2, 2.05, 2, .3, 11, CORAL, True)
textbox(slide, "A focused client for three jobs", 1.2, 2.55, 3.4, .5, 24, WHITE, True, "Aptos Display")
bullet_block(slide, ["Extract a document with OCR or VLM", "Ask a voice question with an image", "Inspect API, Redis, and full health"], 1.2, 3.45, 3.2, 1.6, 16, MINT)
box(slide, 5.55, 1.7, 7.0, 1.4, MINT, radius=True)
textbox(slide, "Thin HTTP client", 5.9, 2.0, 2.5, .3, 19, NAVY, True)
textbox(slide, "api_request() injects Authorization: Bearer <token> and centralizes timeout and error handling.", 5.9, 2.4, 5.9, .35, 15, INK)
box(slide, 5.55, 3.4, 7.0, 1.4, WHITE, line=MINT, radius=True)
textbox(slide, "Useful response surfaces", 5.9, 3.7, 3.2, .3, 19, TEAL, True)
textbox(slide, "Extracted text | structured invoice JSON | transcript | answer | storage references", 5.9, 4.1, 5.9, .35, 15, INK)
box(slide, 5.55, 5.1, 7.0, 1.4, CORAL, radius=True)
textbox(slide, "Important boundary", 5.9, 5.4, 2.5, .3, 19, NAVY, True)
textbox(slide, "The UI does not call provider SDKs. All AI credentials stay server-side.", 5.9, 5.8, 5.9, .35, 15, NAVY)

# 9 deployment
slide = base("Container strategy: build remotely, run simply", "Docker Hub", 9)
box(slide, .8, 1.8, 3.15, 3.9, NAVY, radius=True)
textbox(slide, "GitHub runner", 1.15, 2.2, 2.4, .35, 21, WHITE, True, "Aptos Display")
textbox(slide, "1. checkout\n2. build API image\n3. build UI image\n4. push to Docker Hub\n\nTags: latest + commit SHA", 1.15, 2.95, 2.4, 1.8, 16, MINT)
box(slide, 5.1, 1.8, 3.15, 3.9, TEAL, radius=True)
textbox(slide, "Docker Hub", 5.45, 2.2, 2.4, .35, 21, WHITE, True, "Aptos Display")
textbox(slide, "document-intelligence-api\ndocument-intelligence-ui\n\nImmutable SHA tags make rollback and diagnosis clearer.", 5.45, 2.95, 2.4, 1.8, 16, WHITE)
box(slide, 9.4, 1.8, 3.15, 3.9, CORAL, radius=True)
textbox(slide, "EC2", 9.75, 2.2, 2.4, .35, 21, NAVY, True, "Aptos Display")
textbox(slide, "docker compose pull\ndocker compose up -d --no-build\n\nNo local or EC2 image build. Only runtime memory is needed.", 9.75, 2.95, 2.4, 1.8, 16, NAVY)

# 10 workflow
slide = base("CI/CD workflow", "GitHub Actions", 10)
box(slide, .8, 1.75, 11.7, 4.9, WHITE, line=MINT, radius=True)
textbox(slide, "push to main", 1.25, 2.15, 2.0, .35, 20, NAVY, True)
box(slide, 3.4, 2.25, 1.9, .08, MUTED)
textbox(slide, "publish", 5.55, 2.15, 1.5, .35, 20, TEAL, True)
box(slide, 7.05, 2.25, 1.9, .08, MUTED)
textbox(slide, "deploy", 9.2, 2.15, 1.5, .35, 20, CORAL, True)
textbox(slide, "actions/checkout\ndocker/login-action\ndocker/setup-buildx-action\ndocker/build-push-action", 1.25, 3.05, 3.0, 1.5, 15, INK)
textbox(slide, "Build API + UI\nPush latest + SHA\nRequire Docker Hub secrets", 5.0, 3.05, 2.8, 1.5, 15, INK)
textbox(slide, "SSH to EC2\nSet IMAGE_TAG=github.sha\nPull exact images\nRestart with --no-build", 8.6, 3.05, 2.9, 1.5, 15, INK)
textbox(slide, "Required secrets: DOCKERHUB_USERNAME, DOCKERHUB_TOKEN, EC2_HOST, EC2_USER, EC2_SSH_KEY", 1.25, 5.55, 10.6, .4, 16, NAVY, True)

# 11 ec2
slide = base("First-time EC2 setup", "Operations", 11)
textbox(slide, "The host is a runtime, not a build machine.", .8, 1.65, 7.5, .45, 21, NAVY, True)
steps = [("Install", "Docker + Compose + Git"), ("Clone", "/opt/document-intelligence"), ("Configure", "production .env"), ("Pull", "docker compose pull"), ("Run", "up -d --no-build")]
for i, (head, sub) in enumerate(steps):
    y = 2.35 + i * .72
    box(slide, .9, y, .5, .42, TEAL, radius=True)
    textbox(slide, str(i + 1), 1.02, y + .08, .25, .2, 12, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, head, 1.7, y + .02, 1.8, .25, 17, NAVY, True)
    textbox(slide, sub, 3.8, y + .02, 4.5, .25, 16, INK)
box(slide, 8.75, 2.2, 3.55, 3.4, NAVY, radius=True)
textbox(slide, "Expose carefully", 9.1, 2.6, 2.8, .35, 21, MINT, True, "Aptos Display")
textbox(slide, "Use a reverse proxy and HTTPS for production.\n\nRestrict CORS_ORIGINS to the real UI origin.\n\nKeep API keys in EC2 .env or a secrets manager.", 9.1, 3.35, 2.7, 1.6, 16, WHITE)

# 12 runbook
slide = base("A practical runbook", "Developer workflow", 12)
commands = """LOCAL UI / API
.\\doc-intell\\Scripts\\Activate.ps1
pip install -r requirements.txt
uvicorn app.main:app --reload
streamlit run streamlit_app.py

TESTS
.\\doc-intell\\Scripts\\python.exe -m pytest -q

EC2 RUNTIME
docker compose pull
docker compose up -d --no-build
docker compose ps
curl --fail http://localhost:8000/"""
box(slide, .8, 1.7, 5.85, 4.95, NAVY, radius=True)
textbox(slide, "Commands", 1.15, 2.05, 2, .35, 22, MINT, True, "Aptos Display")
textbox(slide, commands, 1.15, 2.75, 4.9, 3.2, 15, WHITE, font="Consolas")
box(slide, 7.0, 1.7, 5.55, 4.95, MINT, radius=True)
textbox(slide, "When something fails", 7.4, 2.05, 3.5, .35, 22, NAVY, True, "Aptos Display")
bullet_block(slide, ["Check API health before debugging the UI.", "Inspect container logs with docker compose logs -f.", "Confirm .env values and CORS_ORIGINS.", "Check the image tag and Docker Hub visibility.", "Provider smoke tests are separate from mocked CI tests."], 7.4, 2.9, 4.4, 2.4, 16, INK)

# 13 summary
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
box(slide, 0, 0, 13.333, 0.2, CORAL)
textbox(slide, "THE CORE IDEA", .85, 1.0, 3, .35, 12, CORAL, True)
textbox(slide, "Keep intelligence in the API.\nKeep interaction in the UI.\nKeep builds in GitHub.\nKeep runtime on EC2.", .85, 1.6, 8.7, 2.8, 34, WHITE, True, "Aptos Display")
textbox(slide, "The result is a deployable system that does not require a memory-heavy local Docker build.", .9, 5.3, 8.8, .55, 18, MINT)
textbox(slide, "Document Intelligence | walkthrough", 9.6, 6.6, 2.8, .3, 11, MINT, align=PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT)
